import io

import pytest

from src.ingest.document_parser import (
    _line_text_and_size,
    _looks_like_heading,
    parse_document,
    parse_html,
    parse_markdown,
    parse_txt,
)

pytestmark = pytest.mark.unit


def _fake_line(text: str, size: float) -> dict:
    return {"spans": [{"text": text, "size": size}]}


def test_line_text_and_size_joins_spans_and_takes_max_size():
    line = {"spans": [{"text": "Hello ", "size": 10.0}, {"text": "World", "size": 12.0}]}
    text, size = _line_text_and_size(line)
    assert text == "Hello World"
    assert size == 12.0


def test_line_text_strips_embedded_nul_bytes():
    # Found by actually running this against a real arXiv PDF and hitting
    # Postgres's "text fields cannot contain NUL bytes" — see the
    # document_parser.py comment and the ingestion pipeline README entry.
    line = {"spans": [{"text": "bad\x00text", "size": 10.0}]}
    text, _ = _line_text_and_size(line)
    assert "\x00" not in text
    assert text == "badtext"


def test_heading_needs_meaningfully_larger_size():
    assert _looks_like_heading("Introduction", size=11.5, body_size=10.0) is True
    assert _looks_like_heading("Introduction", size=10.2, body_size=10.0) is False


def test_heading_rejects_long_text():
    long_text = "x" * 100
    assert _looks_like_heading(long_text, size=20.0, body_size=10.0) is False


def test_heading_rejects_sentence_trailing_punctuation():
    assert _looks_like_heading("This looks like a sentence.", size=20.0, body_size=10.0) is False


def test_heading_rejects_empty_text():
    assert _looks_like_heading("", size=20.0, body_size=10.0) is False


def test_heading_rejects_zero_body_size():
    assert _looks_like_heading("Short", size=20.0, body_size=0.0) is False


# ---------------------------------------------------------------------
# Phase 3, Stage 1: multi-format parsing. Every parser here must return
# the exact same ParsedDocument/ParsedSection shape parse_pdf already
# does — chunk_document downstream doesn't (and shouldn't) know or care
# which format produced it.
# ---------------------------------------------------------------------


def _build_docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This is the intro paragraph.")
    document.add_heading("Results", level=1)
    document.add_paragraph("This is the results paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Accuracy"
    table.rows[1].cells[1].text = "0.95"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_parse_docx_splits_on_heading_styles_and_extracts_table_content():
    from src.ingest.document_parser import parse_docx

    result = parse_docx(_build_docx_bytes())

    assert [s.heading for s in result.sections] == ["Introduction", "Results"]
    assert "intro paragraph" in result.sections[0].text
    # Table appears under "Results" (the heading it follows in document
    # order) and its content is real extracted text, not skipped.
    assert "Accuracy" in result.sections[1].text
    assert "0.95" in result.sections[1].text


def test_parse_docx_table_between_two_paragraphs_stays_in_document_order():
    # Real gap python-docx's own .paragraphs-only walk has: a table
    # sitting between two paragraphs under the SAME heading must not be
    # silently dropped or reordered — _iter_docx_body's whole reason for
    # existing is walking document.element.body directly instead.
    from docx import Document

    from src.ingest.document_parser import parse_docx

    document = Document()
    document.add_heading("Section", level=1)
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Table content"
    document.add_paragraph("After the table.")
    buf = io.BytesIO()
    document.save(buf)

    result = parse_docx(buf.getvalue())

    assert len(result.sections) == 1
    text = result.sections[0].text
    assert "Before the table" in text
    assert "Table content" in text
    assert "After the table" in text


def test_parse_txt_is_one_section_with_no_heading():
    result = parse_txt(b"Just some plain text.\nA second line.")
    assert len(result.sections) == 1
    assert result.sections[0].heading is None
    assert "Just some plain text." in result.sections[0].text


def test_parse_txt_replaces_bad_bytes_instead_of_failing():
    result = parse_txt(b"good text \xff\xfe bad bytes")
    assert len(result.sections) == 1  # doesn't raise


def test_parse_txt_empty_content_produces_no_sections():
    assert parse_txt(b"   \n  ").sections == []


def test_parse_markdown_splits_on_heading_lines():
    content = b"# Title\nIntro text.\n\n## Subsection\nMore text here.\n"
    result = parse_markdown(content)
    assert [s.heading for s in result.sections] == ["Title", "Subsection"]
    assert "Intro text." in result.sections[0].text
    assert "More text here." in result.sections[1].text


def test_parse_markdown_ignores_non_heading_hashes():
    # A line like "#nothash" (no space after #) must not be mistaken for
    # a heading — matches the regex's explicit \s requirement.
    content = b"#nothash is not a heading\nReal text."
    result = parse_markdown(content)
    assert len(result.sections) == 1
    assert result.sections[0].heading is None


def test_parse_html_strips_script_and_style_and_splits_on_headings():
    content = b"""
    <html><head><style>body{color:red}</style></head>
    <body>
      <h1>Overview</h1>
      <p>First paragraph.</p>
      <script>alert('should not appear')</script>
      <h2>Details</h2>
      <p>Second paragraph.</p>
    </body></html>
    """
    result = parse_html(content)
    assert [s.heading for s in result.sections] == ["Overview", "Details"]
    assert "First paragraph." in result.sections[0].text
    assert "Second paragraph." in result.sections[1].text
    full_text = " ".join(s.text for s in result.sections)
    assert "alert" not in full_text
    assert "color:red" not in full_text


def test_parse_document_dispatches_by_extension():
    result = parse_document("notes.txt", b"hello world")
    assert result.sections[0].text == "hello world"

    result = parse_document("notes.md", b"# H\nbody")
    assert result.sections[0].heading == "H"


def test_parse_document_dispatch_is_case_insensitive():
    result = parse_document("NOTES.TXT", b"hello")
    assert result.sections[0].text == "hello"


def test_parse_document_unsupported_extension_raises_not_implemented():
    with pytest.raises(NotImplementedError):
        parse_document("archive.zip", b"whatever")


# ---------------------------------------------------------------------
# Phase 3, Stage 2: per-page update. parse_pdf must track which page
# each section started on — the exact point this codebase's own earlier
# comments identified as where page info used to be silently discarded.
# ---------------------------------------------------------------------


def _build_multi_page_pdf_bytes() -> bytes:
    # Enough body-text lines (size 11) that statistics.mode reliably
    # picks 11 as body_size over the two size-18 heading lines — matches
    # the real-world padding this project's own live-verification tests
    # needed for the same heuristic (see conversation history: a
    # too-short fixture makes the heading size tie with/beat the body
    # size and the heuristic detects zero headings).
    import pymupdf

    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 100), "Introduction", fontsize=18)
    for i, line in enumerate(["This is the introduction text on page one.", "More body text here.", "Even more body text."]):
        page1.insert_text((72, 140 + i * 20), line, fontsize=11)
    page2 = doc.new_page()
    page2.insert_text((72, 100), "Methods", fontsize=18)
    for i, line in enumerate(["This is the methods text on page two.", "More body text here too.", "Even more body text too."]):
        page2.insert_text((72, 140 + i * 20), line, fontsize=11)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def test_parse_pdf_tracks_which_page_each_section_started_on():
    from src.ingest.document_parser import parse_pdf

    result = parse_pdf(_build_multi_page_pdf_bytes())

    assert [s.heading for s in result.sections] == ["Introduction", "Methods"]
    assert result.sections[0].page == 1
    assert result.sections[1].page == 2


def test_non_pdf_parsers_leave_page_as_none():
    # Deliberate scope decision (schema.py's Chunk.page_number comment):
    # only PDF has a real, stable page concept.
    assert parse_txt(b"hello").sections[0].page is None
    assert parse_markdown(b"# H\nbody").sections[0].page is None
