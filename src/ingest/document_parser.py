from __future__ import annotations

import io
import re
import statistics
from dataclasses import dataclass
from pathlib import Path

import httpx
import pymupdf
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from src.ingest.paper_source import contact_header

_MIN_HEADING_SIZE_RATIO = 1.15
_MAX_HEADING_CHARS = 80


@dataclass(frozen=True)
class ParsedSection:
    heading: str | None
    text: str
    # Phase 3 stage 2: only ever set by parse_pdf (1-indexed) — every
    # other parser leaves this None, since only a PDF has a real, stable
    # page concept (see src/store/schema.py's Chunk.page_number for the
    # full reasoning). Defaulted so parse_docx/parse_txt/parse_markdown/
    # parse_html don't need to know this field exists.
    page: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    sections: list[ParsedSection]


def fetch_pdf_bytes(pdf_url: str) -> bytes:
    response = httpx.get(
        pdf_url,
        headers={"User-Agent": contact_header()},
        timeout=60.0,
        follow_redirects=True,
    )
    response.raise_for_status()
    return response.content


def _line_text_and_size(line: dict) -> tuple[str, float]:
    # PDF text extraction can yield embedded NUL bytes (certain embedded
    # font/ligature encodings decode to \x00) — found by actually running
    # this against a real paper and hitting Postgres's "text fields cannot
    # contain NUL bytes" at the write step. Strip at the source so nothing
    # downstream has to know about it.
    text = "".join(span["text"] for span in line["spans"]).replace("\x00", "").strip()
    size = max((span["size"] for span in line["spans"]), default=0.0)
    return text, size


def _looks_like_heading(text: str, size: float, body_size: float) -> bool:
    if not text or len(text) > _MAX_HEADING_CHARS:
        return False
    if text.endswith((".", ",", ";")):
        return False
    return body_size > 0 and size >= body_size * _MIN_HEADING_SIZE_RATIO


def parse_pdf(pdf_bytes: bytes) -> ParsedDocument:
    """Layout-aware extraction: section headings survive as section
    boundaries rather than being flattened into one text blob (see the
    Infrastructure note on structured parsing).

    Heading detection is a font-size heuristic: a line whose text is
    short, doesn't trail off mid-sentence, and is meaningfully larger
    than the document's own body text size. This is a standard, publicly
    documented PDF-parsing technique — not a specific implementation
    copied from anywhere (see the plan's IP posture note) — deliberately
    simpler than a trained layout model. Revisit if A7's chunking
    ablation shows section boundaries landing wrong.
    """
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    try:
        # Phase 3 stage 2: page.number (0-indexed) is carried alongside
        # each line now — the exact point this project's own earlier
        # comments identified as where page information used to be
        # discarded (flattened into a (text, size)-only tuple with no
        # page index). Stored 1-indexed below since that's what a human
        # (and the per-page update endpoint's URL) actually means by
        # "page 3."
        lines: list[tuple[str, float, int]] = []
        for page in doc:
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:  # 0 = text block; skip images/drawings
                    continue
                for line in block["lines"]:
                    text, size = _line_text_and_size(line)
                    if text:
                        lines.append((text, size, page.number + 1))
    finally:
        doc.close()

    if not lines:
        return ParsedDocument(sections=[])

    body_size = statistics.mode(size for _, size, _ in lines)

    sections: list[ParsedSection] = []
    current_heading: str | None = None
    current_lines: list[str] = []
    # A section's page is the page its FIRST line started on — an
    # approximation for a section that spans a page break, consistent
    # with this chunker's existing section-then-window approach (a chunk
    # already doesn't respect page boundaries within a section either).
    current_page: int | None = None

    for text, size, page_number in lines:
        if _looks_like_heading(text, size, body_size):
            if current_lines:
                sections.append(ParsedSection(heading=current_heading, text=" ".join(current_lines), page=current_page))
            current_heading = text
            current_lines = []
            current_page = page_number
        else:
            if current_page is None:
                current_page = page_number
            current_lines.append(text)

    if current_lines:
        sections.append(ParsedSection(heading=current_heading, text=" ".join(current_lines), page=current_page))

    return ParsedDocument(sections=sections)


def _iter_docx_body(document: DocxDocument):
    """python-docx's own `.paragraphs`/`.tables` don't preserve document
    order or interleave the two — a table between two paragraphs is
    invisible to a `.paragraphs`-only walk. Walking `document.element.body`
    directly and dispatching each child by its real XML tag (`<w:p>` vs
    `<w:tbl>`) is the standard, documented way to get both in the order
    they actually appear — needed here because a table's content must
    land under whichever heading it actually follows, not be silently
    dropped or reordered to the end.
    """
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, document)


def _docx_table_text(table: DocxTable) -> str:
    # Rendered as flowing text, not a structured field on ParsedSection —
    # the chunker only ever consumes plain text, and a "|"-joined row is
    # legible both to a human reading a citation and to the embedding
    # model, without a schema change to carry real table structure
    # through a pipeline that has never needed it for PDFs either.
    rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return "\n".join(row for row in rows if row.strip())


def parse_docx(content: bytes) -> ParsedDocument:
    """DOCX has real heading structure (paragraph styles), unlike PDF's
    font-size heuristic — "Heading 1"/"Heading 2"/etc. styles are exactly
    what Word itself uses to build a table of contents, so trusting them
    directly is more reliable here than PDF's heuristic ever can be.
    Table content is extracted (see _docx_table_text), not skipped —
    real content in a spec-sheet-style upload, not decoration.
    """
    document = DocxDocument(io.BytesIO(content))

    sections: list[ParsedSection] = []
    current_heading: str | None = None
    current_parts: list[str] = []

    for element in _iter_docx_body(document):
        if isinstance(element, DocxParagraph):
            style_name = element.style.name if element.style else ""
            text = element.text.strip()
            if style_name.startswith("Heading") and text:
                if current_parts:
                    sections.append(ParsedSection(heading=current_heading, text="\n".join(current_parts)))
                current_heading = text
                current_parts = []
            elif text:
                current_parts.append(text)
        elif isinstance(element, DocxTable):
            table_text = _docx_table_text(element)
            if table_text:
                current_parts.append(table_text)

    if current_parts:
        sections.append(ParsedSection(heading=current_heading, text="\n".join(current_parts)))

    return ParsedDocument(sections=sections)


def parse_txt(content: bytes) -> ParsedDocument:
    """Plain text has no structure to extract at all — one section, no
    heading. `errors="replace"` rather than failing on a bad byte: a
    single mis-encoded character in an otherwise-fine upload shouldn't
    reject the whole file.
    """
    text = content.decode("utf-8", errors="replace").strip()
    if not text:
        return ParsedDocument(sections=[])
    return ParsedDocument(sections=[ParsedSection(heading=None, text=text)])


_MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def parse_markdown(content: bytes) -> ParsedDocument:
    """A line-prefix heuristic (`^#{1,6}\\s`), not a real Markdown parser —
    deliberately, same honesty as PDF's font-size heuristic: this project
    only ever needs section BOUNDARIES for chunking, not a rendered
    document, so pulling in a markdown-to-HTML renderer just to throw the
    HTML away again would be real, pointless complexity.
    """
    text = content.decode("utf-8", errors="replace")
    sections: list[ParsedSection] = []
    current_heading: str | None = None
    current_lines: list[str] = []

    for raw_line in text.splitlines():
        match = _MARKDOWN_HEADING_RE.match(raw_line)
        if match:
            if current_lines:
                sections.append(ParsedSection(heading=current_heading, text="\n".join(current_lines).strip()))
            current_heading = match.group(2).strip()
            current_lines = []
        elif raw_line.strip():
            current_lines.append(raw_line)

    if current_lines:
        sections.append(ParsedSection(heading=current_heading, text="\n".join(current_lines).strip()))

    return ParsedDocument(sections=[s for s in sections if s.text])


_HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_BLOCK_TAGS = {"p", "li", "td", "th", "blockquote", "pre"}


def parse_html(content: bytes) -> ParsedDocument:
    """Strips script/style (never real document content, and script tags
    especially must never reach the chunker/embedder as if they were
    prose) and walks block-level tags in document order, splitting
    sections on real `<h1>`-`<h6>` markup — actual document structure,
    same reliability class as DOCX's paragraph styles, not a heuristic.
    """
    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()

    sections: list[ParsedSection] = []
    current_heading: str | None = None
    current_parts: list[str] = []

    for tag in soup.find_all(_HEADING_TAGS | _BLOCK_TAGS):
        text = tag.get_text(strip=True)
        if not text:
            continue
        if tag.name in _HEADING_TAGS:
            if current_parts:
                sections.append(ParsedSection(heading=current_heading, text="\n".join(current_parts)))
            current_heading = text
            current_parts = []
        else:
            current_parts.append(text)

    if current_parts:
        sections.append(ParsedSection(heading=current_heading, text="\n".join(current_parts)))

    return ParsedDocument(sections=sections)


_PARSERS_BY_EXTENSION = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".txt": parse_txt,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".html": parse_html,
    ".htm": parse_html,
}


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    """Phase 3: dispatches on file extension to the right format-specific
    parser — every parser above returns the same ParsedDocument shape,
    so nothing downstream of this function (chunk_document, embedding,
    storage) needs to know or care which format a given upload was.
    Raises the same NotImplementedError shape this codebase already uses
    for "unsupported configuration" elsewhere (see embedder.py's
    provider switch) — friendly_error_message already has a branch for
    it, so an unsupported file extension surfaces as a clear in-page
    message, not a raw traceback, with zero new error-handling code.
    """
    ext = Path(filename or "").suffix.lower()
    parser = _PARSERS_BY_EXTENSION.get(ext)
    if parser is None:
        raise NotImplementedError(f"unsupported file type {ext!r} — supported: {sorted(_PARSERS_BY_EXTENSION)}")
    return parser(content)
